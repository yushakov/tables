from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from django.utils import timezone
from django.conf import settings

gDocFiles = ["client_contract.docx",
             "sub_contract.docx"]
gTypesOfWork = ["residential house renovation",
                "HMO conversion",
                "garage conversion",
                "loft conversion",
                "extension",
                "HMO conversion",
                "residential to commercial",
                "commercial to residential",
                "maintenance"]
gDefectPeriods = [6, 12, 24, 36]


def docx_set_cell_border(cell, border_color, border_width):
    # This function adds a border to a cell
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    
    # Set the border on all sides
    for side in ('top', 'left', 'bottom', 'right'):
        tcBorders = OxmlElement('w:tcBorders')
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), str(border_width))
        border.set(qn('w:color'), border_color)
        tcBorders.append(border)
        tcPr.append(tcBorders)


def produce_document(data, construct):
    doc_file_dir = settings.GENDOC_DIR
    file_name = gDocFiles[int(data["docx_file"])]
    full_file_path = doc_file_dir / file_name
    doc = Document(full_file_path)
    now = timezone.now()
    todays_date = now.strftime("%d.%m.%Y")
    replacement = {'#todays date#': todays_date,
                   '#client name#': data['client_name'],
                   '#client address#': data['client_address'],
                   '#project start date#': str(data['start_date']),
                   '#project end date#': str(data['end_date']),
                   '#type of work#': gTypesOfWork[int(data['type_of_work'])], #TODO
                   '#total amount including profit and vat#': data['amount_total_profit_vat'],
                   '#defects period#': str(gDefectPeriods[int(data['defects_period'])]),
                   '#planning consents#': data['planning_consents'],
                   '#party wall consents#': data['party_wall_consents'],
                   '#building regulations#': data['building_regulations'],
                   '#utility water#': data['utility_water'],
                   '#principal designer#': data['principal_designer'],
                   '#principal contractor#': data['principal_contractor'],
                   '#project address#': data['project_address'],
                   '#job description#': data['job_description'],
                   '#deposit percent#': str(construct.deposit_percent_expect),
                   '#deposit#': data['deposit'],
                   }
    for paragraph in doc.paragraphs:
        for key in replacement.keys():
            if key in paragraph.text:
                for run in paragraph.runs:
                    if key in run.text:
                        text = run.text.replace(key, replacement[key])
                        run.text = text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        for key in replacement.keys():
                            if key in run.text:
                                text = run.text.replace(key, replacement[key])
                                run.text = text
    choice_list_key = "#list of works with client prices#"
    choice_list_paragraph = None
    for paragraph in doc.paragraphs:
        if choice_list_key in paragraph.text:
            choice_list_paragraph = paragraph
            break
    if choice_list_paragraph is not None:
        choices = construct.choice_set.all()
        table = doc.add_table(rows=1, cols=7)
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Task name'
        header_cells[1].text = 'Unit price'
        header_cells[2].text = 'Quantity'
        header_cells[3].text = 'Unit'
        header_cells[4].text = 'Total'
        header_cells[5].text = 'Date start'
        header_cells[6].text = 'Planned days'
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        for row in table.rows:
            row.cells[0].width = Inches(2.0)
        total_amount = 0.0
        for choice in choices:
            row_cells = table.add_row().cells
            row_cells[0].text = choice.name_txt
            row_cells[1].text = "£ " + str(construct.with_all_profits_and_vat(choice.price_num))
            row_cells[2].text = str(choice.quantity_num)
            row_cells[3].text = str(choice.units_of_measure_text)
            full_work_price = construct.with_all_profits_and_vat(choice.quantity_num * choice.price_num)
            total_amount += full_work_price
            row_cells[4].text = "£" + str(round(full_work_price, 2))
            row_cells[5].text = str(choice.plan_start_date)
            row_cells[6].text = str(choice.plan_days_num)
        for row in table.rows:
            for cell in row.cells:
                docx_set_cell_border(cell, 'D9D9D9', 6)
        row_cells = table.add_row().cells
        row_cells[3].text = "Total:"
        row_cells[4].text = "£ " + str(round(total_amount, 2))
        row_cells[3].paragraphs[0].runs[0].bold = True
        choice_list_paragraph.text = ""

    doc.save(doc_file_dir / "new_document.docx")